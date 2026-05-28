"""实验数据加载器 —— h-BN 1T1M 忆阻器 I-V 测量数据。

数据来源: Nature 2023 (Zhu, Pazos et al.)
文件:
  1T1M写入.xlsx — 53 个 Set 周期 (S1-S53), 0→+5V→0
  1T1M擦除.xlsx — 53 个 Reset 周期 (R1-R53), 0→-1.7V→0

每周期 503 点 (Set) 或 173 点 (Reset), 列: voltage, current
"""

from __future__ import annotations
import numpy as np
from pathlib import Path
import openpyxl


class ExperimentData:
    """单个 I-V 扫描周期的数据容器。

    属性:
        voltage: 电压数组 [V]
        current: 电流数组 [A]
        current_ma: 电流 [mA] (便捷单位转换)
        current_ua: 电流 [µA] (便捷单位转换)
        v_range: (V_min, V_max) 范围
        i_max: 最大 |电流|
    """

    def __init__(self, voltage: np.ndarray, current: np.ndarray,
                 label: str = "", cycle_id: int = 0):
        self.voltage = voltage
        self.current = current
        self.label = label              # e.g. "set_S5"
        self.cycle_id = cycle_id        # 周期编号 (1-53)

    @property
    def current_ma(self) -> np.ndarray:
        """电流 [mA]"""
        return self.current * 1e3

    @property
    def current_ua(self) -> np.ndarray:
        """电流 [µA]"""
        return self.current * 1e6

    @property
    def v_range(self) -> tuple[float, float]:
        """电压范围 (min, max) [V]"""
        return float(self.voltage.min()), float(self.voltage.max())

    @property
    def i_max(self) -> float:
        """最大 |电流| [A]"""
        return float(abs(self.current).max())


def load_nature_dataset(data_dir: str | Path = None) -> dict:
    """加载完整 Nature h-BN 1T1M 数据集。

    Returns:
        dict {
            'set':      53 个 ExperimentData (Set/Write 周期)
            'reset':    53 个 ExperimentData (Reset/Erase 周期)
            'set_mean': 所有 Set 周期的平均 I-V (ExperimentData)
            'set_std':  所有 Set 周期的 I 标准差 (ExperimentData)
        }
    """
    if data_dir is None:
        data_dir = Path(__file__).parent.parent.parent / "experiment data" / "nature h-BN"

    data_dir = Path(data_dir)
    set_cycles = _load_xlsx(data_dir / "1T1M写入.xlsx", prefix="S", sweep_type="set")
    reset_cycles = _load_xlsx(data_dir / "1T1M擦除.xlsx", prefix="R", sweep_type="reset")

    # Compute mean and std across set cycles
    set_mean, set_std = _compute_cycle_stats(set_cycles)

    return {
        "set": set_cycles,
        "reset": reset_cycles,
        "set_mean": set_mean,
        "set_std": set_std,
    }


def _load_xlsx(path: str | Path, prefix: str, sweep_type: str) -> list[ExperimentData]:
    """Load all sheets from one xlsx file."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Data file not found: {path}")

    wb = openpyxl.load_workbook(str(path), data_only=True)
    cycles = []

    for sn in wb.sheetnames:
        ws = wb[sn]
        n_rows = ws.max_row - 1  # Skip header

        v = np.zeros(n_rows)
        i = np.zeros(n_rows)
        for r in range(2, ws.max_row + 1):
            v[r - 2] = float(ws.cell(r, 1).value or 0)
            i[r - 2] = float(ws.cell(r, 2).value or 0)

        cycle_num = int(sn[1:])  # e.g., "S5" -> 5
        cycles.append(ExperimentData(v, i, label=f"{sweep_type}_{sn}", cycle_id=cycle_num))

    return sorted(cycles, key=lambda c: c.cycle_id)


def _compute_cycle_stats(cycles: list[ExperimentData]) -> tuple[ExperimentData | None, ExperimentData | None]:
    """Compute mean and std I-V across cycles (assumes same V grid)."""
    if not cycles:
        return None, None

    # Use the voltage grid from the first cycle
    v_ref = cycles[0].voltage
    n_pts = len(v_ref)
    all_i = np.zeros((len(cycles), n_pts))

    for idx, c in enumerate(cycles):
        if len(c.voltage) == n_pts:
            all_i[idx, :] = c.current
        else:
            all_i[idx, :] = np.interp(v_ref, c.voltage, c.current)

    i_mean = np.mean(all_i, axis=0)
    i_std = np.std(all_i, axis=0)

    mean_data = ExperimentData(v_ref, i_mean, label="set_mean", cycle_id=0)
    std_data = ExperimentData(v_ref, i_std, label="set_std", cycle_id=0)

    return mean_data, std_data
